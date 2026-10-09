import csv
import html
from pathlib import Path
from typing import List, Dict, Any
import openpyxl

COLUMNS = [
    "Nom de la ferme",
    "Propriétaire",
    "Région / Province",
    "Pays",
    "Superficie (ha ou m²)",
    "Type de culture",
    "Courriel",
    "Téléphone",
    "Site Web",
    "Adresse",
    "Ville",
    "Code Postal",
    "Certifications",
    "Statut",
    "Notes"
]

def clean_value(val: Any) -> str:
    if val is None:
        return ""
    text = html.unescape(str(val)).strip()
    return " ".join(text.split()) if "\n" not in text else text.strip()

def export_to_csv(records: List[Dict[str, Any]], filepath: Path):
    filepath.parent.mkdir(parents=True, exist_ok=True)
    with open(filepath, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for r in records:
            clean_row = {col: clean_value(r.get(col, "")) for col in COLUMNS}
            writer.writerow(clean_row)

def export_to_excel(records: List[Dict[str, Any]], filepath: Path):
    filepath.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fermes & Maraîchers"
    
    # Headers
    ws.append(COLUMNS)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    
    # Rows
    for r in records:
        row_vals = [clean_value(r.get(col, "")) for col in COLUMNS]
        ws.append(row_vals)
        
    wb.save(filepath)
